#!/usr/bin/env python3
"""Generate ScoutMatch AI final submission presentation (6 main + 5 appendix slides)."""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT_PPTX = HERE / "ScoutMatch_AI_Final_Submission_Presentation.pptx"
OUT_PDF = HERE / "ScoutMatch_AI_Final_Submission_Presentation.pdf"
PUBLIC_URL = "http://3.239.47.249/"
V1_TOOL_NAMES = ["PlanMatchTactics", "SubmitPlayerSelection", "FinalizeCurrentLineup", "GenerateLineupBoard"]

# Theme
NAVY = RGBColor(0x0B, 0x1F, 0x33)
NAVY2 = RGBColor(0x12, 0x2B, 0x45)
GREEN = RGBColor(0x2E, 0xCC, 0x71)
CYAN = RGBColor(0x00, 0xD4, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xC8, 0xD6, 0xE5)
MUTED = RGBColor(0x8E, 0xA6, 0xBD)
GOLD = RGBColor(0xF3, 0x9C, 0x12)

ASSETS = {
    "homepage": ROOT / "artifacts/evidence/ui_polish_hero_clean_manual_final_1366x768.png",
    "rag": ROOT / "artifacts/evidence/business_workflow_v2_staging/03_grounded_squad_analysis.png",
    "confirm": ROOT / "artifacts/evidence/business_workflow_v2_staging/09_critical_decision_confirm_card.png",
    "budget": ROOT / "artifacts/evidence/business_workflow_v2_staging/10_critical_decision_result.png",
    "lineup": ROOT / "artifacts/evidence/business_workflow_v2_staging/12_visual_squad_board.png",
}


def _font(size: int, bold: bool = False):
    names = ["segoeui.ttf", "arial.ttf", "calibri.ttf"]
    for name in names:
        try:
            return ImageFont.truetype(name, size=size)
        except OSError:
            continue
    return ImageFont.load_default()



def _set_slide_bg(slide, color: RGBColor = NAVY) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_footer(slide, text: str = "ScoutMatch AI  |  Avidan Mendelman") -> None:
    box = slide.shapes.add_textbox(Inches(0.4), Inches(7.0), Inches(9.2), Inches(0.35))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(11)
    p.font.color.rgb = MUTED


def _add_title_block(slide, title: str, subtitle: str = "", appendix: bool = False) -> None:
    accent = GOLD if appendix else GREEN
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.45), Inches(0.35), Inches(0.08), Inches(0.85))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    tbox = slide.shapes.add_textbox(Inches(0.65), Inches(0.32), Inches(8.8), Inches(0.7))
    tp = tbox.text_frame.paragraphs[0]
    prefix = "APPENDIX — " if appendix else ""
    tp.text = prefix + title
    tp.font.size = Pt(30 if appendix else 34)
    tp.font.bold = True
    tp.font.color.rgb = WHITE
    if subtitle:
        sbox = slide.shapes.add_textbox(Inches(0.65), Inches(0.95), Inches(8.8), Inches(0.45))
        sp = sbox.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(18)
        sp.font.color.rgb = CYAN


def _add_bullets(slide, items: list[str], left: float = 0.65, top: float = 1.45, width: float = 8.7, height: float = 5.2, size: int = 20) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = LIGHT
        p.space_after = Pt(8)


def _add_image_fit(slide, path: Path, left: float, top: float, max_w: float, max_h: float, caption: str = "") -> None:
    if not path.exists():
        return
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    w = iw * ratio
    h = ih * ratio
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(w), height=Inches(h))
    if caption:
        cbox = slide.shapes.add_textbox(Inches(left), Inches(top + h + 0.05), Inches(max_w), Inches(0.3))
        cp = cbox.text_frame.paragraphs[0]
        cp.text = caption
        cp.font.size = Pt(11)
        cp.font.color.rgb = MUTED


def _draw_architecture(slide) -> None:
    boxes = [
        (0.5, 2.0, 1.35, 0.65, "Browser", GREEN),
        (2.1, 2.0, 1.35, 0.65, "EC2 :80", CYAN),
        (3.7, 2.0, 1.35, 0.65, "Docker", CYAN),
        (5.3, 2.0, 1.35, 0.65, "Flask", CYAN),
        (6.9, 2.0, 1.35, 0.65, "boto3 + IAM", GREEN),
        (4.0, 3.2, 2.2, 0.75, "Bedrock Agent", GREEN),
        (2.0, 4.5, 1.7, 0.6, "Guardrail", GOLD),
        (4.0, 4.5, 1.7, 0.6, "Knowledge Base", CYAN),
        (6.0, 4.5, 1.7, 0.6, "4 Action Groups", GREEN),
        (2.0, 5.7, 1.7, 0.6, "4 Lambdas", GREEN),
        (4.0, 5.7, 1.7, 0.6, "DynamoDB", CYAN),
        (6.0, 5.7, 1.7, 0.6, "S3 docs + SVG", CYAN),
        (4.0, 6.55, 2.2, 0.55, "Flask SVG proxy", GREEN),
    ]
    for x, y, w, h, label, color in boxes:
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY2
        shape.line.color.rgb = color
        shape.line.width = Pt(2)
        tf = shape.text_frame
        tf.text = label
        tf.paragraphs[0].font.size = Pt(13)
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def _draw_flow(slide) -> None:
    steps = [
        "Scout asks",
        "Agent + KB",
        "Tool if needed",
        "Confirm / Deny",
        "DynamoDB",
        "Lineup board",
    ]
    x = 0.45
    for i, step in enumerate(steps):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.2), Inches(1.35), Inches(0.7))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY2
        shape.line.color.rgb = GREEN if i % 2 == 0 else CYAN
        shape.line.width = Pt(2)
        tf = shape.text_frame
        tf.text = step
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        if i < len(steps) - 1:
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + 1.38), Inches(2.55), Inches(x + 1.55), Inches(2.55))
            conn.line.color.rgb = MUTED
            conn.line.width = Pt(2)
        x += 1.55


def build_pptx() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    _set_slide_bg(s)
    _add_title_block(s, "About Me")
    box = s.shapes.add_textbox(Inches(0.65), Inches(1.55), Inches(8.5), Inches(4.8))
    tf = box.text_frame
    tf.word_wrap = True
    lines = [
        ("Avidan Mendelman", 28, WHITE, True),
        ("The Open University of Israel", 22, LIGHT, False),
        ("B.Sc. in Computer Science", 22, LIGHT, False),
        ("", 10, LIGHT, False),
        ("I am passionate about software development and building useful AI applications.", 18, CYAN, False),
    ]
    for i, (txt, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.space_after = Pt(8)
    s.notes_slide.notes_text_frame.text = "Introduce yourself (~45s). Open University CS; passion for practical AI."
    _add_footer(s)

    s = prs.slides.add_slide(blank)
    _set_slide_bg(s)
    _add_title_block(s, "ScoutMatch AI — Evidence-Based Football Recruitment")
    _add_bullets(
        s,
        [
            "Club finished fourth last season — strengthen before the transfer window closes.",
            "User: football scout / recruitment analyst.",
            "Analyzes club documents, squad data, candidate reports, and operational decisions.",
            "Identify weaknesses, transfer-out reviews, scouting missions, critical recommendations, budget on Confirm, visual lineup.",
            f"Live demo: {PUBLIC_URL}",
        ],
        top=1.35,
        width=4.35,
        height=5.0,
        size=15,
    )
    _add_image_fit(s, ASSETS["homepage"], 5.05, 1.35, 4.45, 5.0, "Polished homepage — six quick-start cards")
    s.notes_slide.notes_text_frame.text = "Opening-season recruitment scenario (~60s). Evidence-based; human approval."
    _add_footer(s)

    s = prs.slides.add_slide(blank)
    _set_slide_bg(s)
    _add_title_block(s, "Technologies Used")
    groups = [
        ("Application", ["Python", "Flask", "HTML / CSS / JavaScript"]),
        ("Deployment", ["Docker", "Amazon EC2", "Git / GitHub"]),
        ("AWS AI", ["Amazon Bedrock Agent", "Amazon Bedrock Knowledge Base", "RAG", "Guardrail", "boto3", "IAM Role"]),
        ("AWS Services", ["Amazon S3", "AWS Lambda", "Amazon DynamoDB"]),
        ("Tools", ["Exactly 4 Tools", "Exactly 4 Action Groups", "Exactly 4 Lambda Functions"]),
    ]
    x = 0.45
    for label, items in groups:
        shape = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.45), Inches(1.55), Inches(0.42))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY2
        shape.line.color.rgb = GOLD
        tf = shape.text_frame
        tf.text = label
        tf.paragraphs[0].font.size = Pt(11)
        tf.paragraphs[0].font.color.rgb = GOLD
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        yy = 1.95
        for item in items:
            sh = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(yy), Inches(1.55), Inches(0.34))
            sh.fill.solid()
            sh.fill.fore_color.rgb = NAVY2
            sh.line.color.rgb = CYAN
            stf = sh.text_frame
            stf.text = item
            stf.paragraphs[0].font.size = Pt(9)
            stf.paragraphs[0].font.color.rgb = LIGHT
            stf.paragraphs[0].alignment = PP_ALIGN.CENTER
            stf.vertical_anchor = MSO_ANCHOR.MIDDLE
            yy += 0.38
        x += 1.72
    note = s.shapes.add_textbox(Inches(0.45), Inches(6.05), Inches(9.1), Inches(0.7))
    np = note.text_frame.paragraphs[0]
    np.text = "The project uses controlled Tools through Amazon Bedrock Agent Action Groups and AWS Lambda functions. It does not use a standalone MCP server."
    np.font.size = Pt(13)
    np.font.color.rgb = CYAN
    s.notes_slide.notes_text_frame.text = "Stack overview (~50s). Four tools; no MCP server."
    _add_footer(s)

    s = prs.slides.add_slide(blank)
    _set_slide_bg(s)
    _add_title_block(s, "System Architecture")
    _draw_architecture(s)
    gate = s.shapes.add_textbox(Inches(0.45), Inches(6.05), Inches(9.1), Inches(0.55))
    gp = gate.text_frame.paragraphs[0]
    gp.text = "Guardrail, Confirm / Deny, and human approval boundaries before writes"
    gp.font.size = Pt(14)
    gp.font.color.rgb = GREEN
    s.notes_slide.notes_text_frame.text = "End-to-end flow through EC2, Docker, Flask, Bedrock Agent, KB, Lambdas, DynamoDB, S3, Flask proxy (~60s)."
    _add_footer(s)

    s = prs.slides.add_slide(blank)
    _set_slide_bg(s)
    _add_title_block(s, "Live Demo")
    _add_bullets(
        s,
        [
            "1 Ask: squad priority before transfer window closes.",
            "2 Show: RAG-grounded answer with club-document evidence.",
            "3 Ask: submit Ron Ben Ari recommendation, then Confirm / Deny card.",
            "4 Confirm: 43,000 EUR reserved, 57,000 EUR remaining, Pending management approval.",
            "5 Show: updated proposed lineup and squad-risk board.",
        ],
        top=1.28,
        width=4.2,
        height=3.0,
        size=14,
    )
    _add_image_fit(s, ASSETS["rag"], 0.45, 4.55, 2.85, 1.55, "Grounded analysis")
    _add_image_fit(s, ASSETS["confirm"], 3.45, 4.55, 2.85, 1.55, "Confirm / Deny")
    _add_image_fit(s, ASSETS["lineup"], 6.45, 4.55, 2.85, 1.55, "Lineup board")
    s.notes_slide.notes_text_frame.text = "Demo story (~90s). Live site or screenshots as backup."
    _add_footer(s)

    s = prs.slides.add_slide(blank)
    _set_slide_bg(s)
    _add_title_block(s, "Challenges and Next Steps")
    sections = [
        ("What Works", GREEN, [
            "Public Docker on EC2",
            "RAG answers with sources",
            "Bedrock Agent with 4 controlled tools",
            "Confirm / Deny before writes",
            "DynamoDB state updates",
            "Private S3 lineup SVG via Flask Proxy",
            "Guardrail protection",
        ]),
        ("Technical Challenges", GOLD, [
            "Bedrock returnControl to the website",
            "Idempotency against duplicate writes",
            "Private lineup assets in S3",
            "Human approval boundaries",
        ]),
        ("Next Steps", CYAN, [
            "Real SES email delivery",
            "Full Google Calendar integration",
            "Management approval dashboard",
            "Better player-ranking analytics",
            "Mobile UI improvements",
        ]),
    ]
    x = 0.45
    for title, accent, items in sections:
        hdr = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.35), Inches(2.85), Inches(0.45))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = NAVY2
        hdr.line.color.rgb = accent
        htf = hdr.text_frame
        htf.text = title
        htf.paragraphs[0].font.size = Pt(13)
        htf.paragraphs[0].font.color.rgb = accent
        htf.paragraphs[0].alignment = PP_ALIGN.CENTER
        htf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _add_bullets(s, [f"• {it}" for it in items], left=x, top=1.95, width=2.85, height=3.2, size=13)
        x += 3.05
    stats = [("481", "Tests passed"), ("0", "Failed tests"), ("3/3", "Public browser reliability")]
    sx = 0.45
    for val, label in stats:
        card = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(sx), Inches(5.55), Inches(2.85), Inches(0.85))
        card.fill.solid()
        card.fill.fore_color.rgb = NAVY2
        card.line.color.rgb = GREEN
        ctf = card.text_frame
        ctf.clear()
        p1 = ctf.paragraphs[0]
        p1.text = val
        p1.font.size = Pt(26)
        p1.font.bold = True
        p1.font.color.rgb = GREEN
        p1.alignment = PP_ALIGN.CENTER
        p2 = ctf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(12)
        p2.font.color.rgb = LIGHT
        p2.alignment = PP_ALIGN.CENTER
        sx += 3.05
    url = s.shapes.add_textbox(Inches(0.45), Inches(6.55), Inches(9.1), Inches(0.35))
    up = url.text_frame.paragraphs[0]
    up.text = f"Public URL: {PUBLIC_URL}"
    up.font.size = Pt(14)
    up.font.color.rgb = GOLD
    s.notes_slide.notes_text_frame.text = "Wins, challenges, roadmap (~60s). 481 passed, 0 failed, 3/3 reliability."
    _add_footer(s)

    appendix = [
        ("Exact Four Tools and Lambda Mapping", [
            "1 SubmitCriticalDecisionAndSendEmail — Ron Ben Ari, 43,000 EUR reserved, 57,000 remaining, Pending management approval.",
            "2 OpenTransferOutReviewCase — Daniel Cohen transfer-out, 25,000 EUR estimated release, NOT sold automatically.",
            "3 CreateAndReviewScoutingMission — Ron Ben Ari mission, ICS download and synthetic report, no full Google Calendar API.",
            "4 GenerateVisualSquadAndLineupBoard — 4-3-3 board in private S3 SVG via Flask Proxy, Pending head-coach review.",
            "Exactly 4 Tools, 4 Action Groups, 4 Lambdas, one-to-one mapping. No fifth tool.",
        ]),
        ("Real Business Scenario", [
            "Opening season: club finished fourth; transfer window closing.",
            "Scout analyzes squad gaps and compares Ron Ben Ari within budget.",
            "Ron Ben Ari is NOT automatically signed.",
            "Daniel Cohen is NOT automatically sold.",
            "Lineup is NOT automatically approved — head coach decides.",
        ]),
        ("Knowledge Base vs DynamoDB", [
            "Knowledge Base (S3 + Bedrock): static club documents — budget, tactics, policies, CVs.",
            "DynamoDB: live operational state — reservations, cases, missions, lineup records.",
            "RAG reads KB; tools write DynamoDB only after Confirm.",
        ]),
        ("MCP Clarification", [
            "Uses Amazon Bedrock Agent Action Groups backed by AWS Lambda.",
            "Does NOT implement a standalone MCP server.",
            "Four controlled tools — not open-ended plugin hosting.",
        ]),
        ("Instructor Q and A", [
            "Does it sign players automatically? No — pending management approval only.",
            "Is SNS required? No — optional future extension.",
            "How many tools? Exactly four Action Groups and four Lambdas.",
            f"Public demo: {PUBLIC_URL}",
        ]),
    ]
    for title, bullets in appendix:
        s = prs.slides.add_slide(blank)
        _set_slide_bg(s, NAVY2)
        _add_title_block(s, title, appendix=True)
        _add_bullets(s, bullets, top=1.55, size=16)
        _add_footer(s, "APPENDIX — Presenter: Avidan Mendelman")

    prs.save(OUT_PPTX)
    return prs


def export_pdf() -> bool:
    try:
        import comtypes.client  # type: ignore

        powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
        powerpoint.Visible = 1
        presentation = powerpoint.Presentations.Open(str(OUT_PPTX.resolve()), WithWindow=False)
        presentation.SaveAs(str(OUT_PDF.resolve()), 32)  # ppSaveAsPDF
        presentation.Close()
        powerpoint.Quit()
        return OUT_PDF.exists()
    except Exception as exc:
        print(f"PDF export via PowerPoint COM failed: {exc}")
        return False


def inspect_outputs() -> dict:
    result = {
        "pptx": False,
        "pdf": False,
        "slide_count": 0,
        "main_count": 0,
        "appendix_count": 0,
        "issues": [],
        "sensitive": [],
    }
    if not OUT_PPTX.exists():
        result["issues"].append("PPTX missing")
        return result
    prs = Presentation(str(OUT_PPTX))
    result["slide_count"] = len(prs.slides)
    result["main_count"] = min(6, result["slide_count"])
    result["appendix_count"] = max(0, result["slide_count"] - 6)
    bad = ("arn:aws", "881490130721", ".env", "@gmail", "AKIA", "pem", "token")
    deck_blob = ""
    for i, slide in enumerate(prs.slides, start=1):
        slide_text = ""
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs) or p.text
                slide_text += txt + " "
                low = txt.lower()
                if any(x in low for x in bad):
                    result["sensitive"].append(f"slide {i}: {txt[:50]}")
                if len(txt) > 240:
                    result["issues"].append(f"Long text block slide {i} ({len(txt)} chars)")
        deck_blob += slide_text + " "
        for v1 in V1_TOOL_NAMES:
            if v1 in slide_text:
                result["issues"].append(f"V1 tool name on slide {i}: {v1}")
    if result["slide_count"] != 11:
        result["issues"].append(f"expected 11 slides, got {result['slide_count']}")
    if "3.239.47.249" not in deck_blob:
        result["issues"].append("public URL missing")
    if "481" not in deck_blob:
        result["issues"].append("481 tests count missing")
    if "0" not in deck_blob or "Failed tests" not in deck_blob:
        result["issues"].append("failed test count missing")
    for tool in [
        "SubmitCriticalDecisionAndSendEmail",
        "OpenTransferOutReviewCase",
        "CreateAndReviewScoutingMission",
        "GenerateVisualSquadAndLineupBoard",
    ]:
        if tool not in deck_blob:
            result["issues"].append(f"missing tool: {tool}")
    result["pptx"] = result["slide_count"] == 11 and not [x for x in result["issues"] if not x.startswith("Long")]
    if OUT_PDF.exists() and OUT_PDF.stat().st_size > 10000:
        result["pdf"] = True
    else:
        result["issues"].append("PDF missing or too small")
    return result
    prs = Presentation(str(OUT_PPTX))
    result["slide_count"] = len(prs.slides)
    result["pptx"] = result["slide_count"] == 17
    for i, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs) or p.text
                if any(x in txt for x in ("arn:aws", "881490130721", ".env", "@", "AKIA")):
                    result["issues"].append(f"Sensitive text on slide {i}: {txt[:40]}")
                if len(txt) > 220:
                    result["issues"].append(f"Long text block slide {i} ({len(txt)} chars)")
    if OUT_PDF.exists() and OUT_PDF.stat().st_size > 10000:
        result["pdf"] = True
    else:
        result["issues"].append("PDF missing or too small")
    return result


def write_markdown_files() -> None:
    (HERE / "Speaker_Notes.md").write_text("""# ScoutMatch AI — Speaker Notes

Total main deck time: **5–7 minutes** (appendix for Q&A only).

## Slide 1 — About Me (~45s)
- Avidan Mendelman, The Open University of Israel, B.Sc. Computer Science.
- Passionate about software development and useful AI applications.

## Slide 2 — Project Overview (~60s)
- Club finished fourth; strengthen before transfer window closes.
- Scout/recruitment analyst using club documents, squad data, candidate reports.
- Weaknesses, transfer-out, missions, recommendations, budget on Confirm, visual lineup.
- Live demo: http://3.239.47.249/

## Slide 3 — Technologies Used (~50s)
- Python, Flask, HTML/CSS/JS; Docker on EC2; Git/GitHub.
- Bedrock Agent, Knowledge Base, RAG, Guardrail, boto3, IAM Role; S3, Lambda, DynamoDB.
- Exactly 4 Tools, 4 Action Groups, 4 Lambdas — not a standalone MCP server.

## Slide 4 — System Architecture (~60s)
- Browser to EC2 to Docker to Flask to boto3/IAM to Bedrock Agent to KB/RAG to Action Group to Lambda to DynamoDB or private S3 to Flask Proxy to Browser.
- Guardrail; Confirm/Deny preserves human approval.

## Slide 5 — Live Demo (~90s)
- Squad question, grounded answer, Ron recommendation, Confirm, budget, lineup board.

## Slide 6 — Challenges and Next Steps (~60s)
- What works, technical challenges, next steps.
- 481 tests passed, 0 failed, public browser reliability 3/3.

## Appendix (Q&A only)
- Use A1–A5 only when asked.
""", encoding="utf-8")

    (HERE / "Live_Demo_Script.md").write_text("""# ScoutMatch AI — Live Demo Script

Duration: **2–3 minutes**.

## Prompt 1
```
We finished fourth last season. Analyze our current squad before the transfer window closes. Which position should we prioritize?
```
**Expected:** RAG-grounded answer with source cards.
**Fallback:** Retrieval usually takes 10–20 seconds.
**Screenshot:** artifacts/evidence/business_workflow_v2_staging/03_grounded_squad_analysis.png

## Prompt 2
```
I choose Ron Ben Ari as our right-back candidate. Submit the recommendation for management review.
```
**Expected:** Confirm / Deny card.
**Screenshot:** artifacts/evidence/business_workflow_v2_staging/09_critical_decision_confirm_card.png

## Click Confirm
**Expected:** 43,000 EUR reserved, 57,000 EUR remaining, Pending management approval.
**Screenshot:** artifacts/evidence/business_workflow_v2_staging/10_critical_decision_result.png

## Prompt 3
```
Show me the updated proposed lineup and squad-risk board.
```
**Expected:** Inline 11-player lineup board.
**Screenshot:** artifacts/evidence/business_workflow_v2_staging/12_visual_squad_board.png

## Human approval (say aloud)
Ron is not signed; Daniel Cohen is not sold; lineup is not approved.

## If live demo fails
Show screenshots above and state http://3.239.47.249/
""", encoding="utf-8")

    (HERE / "Presentation_Outline.md").write_text("""# ScoutMatch AI — Presentation Outline

## Main deck (6 slides)
1. About Me
2. ScoutMatch AI — Evidence-Based Football Recruitment
3. Technologies Used
4. System Architecture
5. Live Demo
6. Challenges and Next Steps

## Appendix (Q&A, 5 slides)
A1 Exact Four Tools and Lambda Mapping
A2 Real Business Scenario
A3 Knowledge Base vs DynamoDB
A4 MCP Clarification
A5 Instructor Q&A
""", encoding="utf-8")
def main() -> int:
    build_pptx()
    write_markdown_files()
    pdf_ok = export_pdf()
    inspection = inspect_outputs()
    print(f"PPTX={OUT_PPTX} slides={inspection['slide_count']}")
    print(f"PDF={OUT_PDF} ok={pdf_ok}")
    if inspection["issues"]:
        print("ISSUES:")
        for issue in inspection["issues"]:
            print(f"  - {issue}")
        return 1 if not pdf_ok else 0
    return 0


if __name__ == "__main__":
    sys.exit(main())
